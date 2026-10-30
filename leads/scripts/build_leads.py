#!/usr/bin/env python3
"""Merge sources, enrich generic contacts, write CSV + leads-nn.xlsx."""

from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from enrich_sites import enrich_rows
from ingest_2gis import ingest as ingest_2gis
from nn_common import (
    CSV_FIELDS,
    RAW_DIR,
    ROOT,
    fold,
    in_scope_company,
    in_scope_kindergarten,
    in_scope_restaurant,
    sort_rows,
    upsert,
    write_csv,
)

SEGMENTS = {
    "рестораны": ROOT / "restaurants-nn.csv",
    "детские сады и центры": ROOT / "kindergartens-nn.csv",
    "компании": ROOT / "companies-nn.csv",
}
SHEET_NAMES = {
    "рестораны": "Рестораны",
    "детские сады и центры": "Детские сады и центры",
    "компании": "Компании",
}


def load_json_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        return [x for x in data if isinstance(x, dict) and x.get("название")]
    return []


def collect_all() -> dict[str, dict]:
    index: dict[str, dict] = {}
    for path in sorted(RAW_DIR.glob("*.json")):
        if path.name.startswith("2gis-"):
            continue
        for rec in load_json_rows(path):
            upsert(index, rec)
    gis_paths = sorted(RAW_DIR.glob("2gis-*.json"))
    if gis_paths:
        for rec in ingest_2gis(gis_paths):
            upsert(index, rec)
    return {k: v for k, v in index.items() if keep_record(v)}


def keep_record(rec: dict) -> bool:
    name = rec.get("название") or ""
    sub = rec.get("подсегмент") or ""
    extra = f"{rec.get('источник') or ''} {rec.get('источник_url') or ''}"
    seg = rec.get("сегмент")
    if seg == "рестораны":
        return in_scope_restaurant(name, sub, extra)
    if seg == "детские сады и центры":
        return in_scope_kindergarten(name, sub, extra)
    if seg == "компании":
        if not in_scope_company(name, sub, extra):
            return False
        src = rec.get("источник") or ""
        addr = rec.get("адрес") or ""
        # Firmoteka rating is oblast-wide; keep a row only if maps gave a city street
        # or the address already contains a street/house.
        if src.strip() == "Firmoteka / открытый рейтинг ФНС":
            if not re.search(r"(ул|улица|пр|проспект|пл\.|площадь|шоссе|набереж|\d)", addr, re.I):
                return False
        return True
    return False


def collapse_bank_branches(rows: list[dict]) -> list[dict]:
    """Keep one office per bank brand per district; other companies untouched."""
    bank_seen: set[str] = set()
    out = []
    for rec in rows:
        blob = fold(rec.get("название") or "")
        if rec.get("сегмент") != "компании":
            out.append(rec)
            continue
        brand = ""
        for b in ("сбербанк", "сбер банк", "втб", "альфа-банк", "альфа банк", "тинькофф", "т-банк", "газпромбанк", "россельхозбанк", "открытие", "совкомбанк", "райффайзен", "почта банк", "мтс банк"):
            if b in blob:
                brand = b
                break
        if not brand:
            out.append(rec)
            continue
        key = f"{brand}|{fold(rec.get('район') or '')}"
        if key in bank_seen:
            continue
        bank_seen.add(key)
        out.append(rec)
    return out


def write_xlsx(by_seg: dict[str, list[dict]], path: Path) -> None:
    wb = Workbook()
    # summary first
    ws = wb.active
    ws.title = "Сводка"
    ws.append(["Сегмент", "Записей", "С телефоном", "С сайтом", "С email", "С VK", "Приоритет 1–2"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    for seg, rows in by_seg.items():
        ws.append(
            [
                seg,
                len(rows),
                sum(bool(r.get("телефон")) for r in rows),
                sum(bool(r.get("сайт")) for r in rows),
                sum(bool(r.get("email")) for r in rows),
                sum(bool(r.get("vk")) for r in rows),
                sum(str(r.get("приоритет")) in {"1", "2"} for r in rows),
            ]
        )
    ws.append([])
    ws.append(["Колонка приоритета: 1 — крупные/премиальные, 5 — ниже."])
    ws.append(["Контакты только общие (info@, zakaz@, dsN_nn@mail.52gov.ru). ФИО не собирались."])
    ws.append(["Город: Нижний Новгород, без области."])

    for seg, rows in by_seg.items():
        sheet = wb.create_sheet(SHEET_NAMES[seg])
        sheet.append(CSV_FIELDS)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        fills = {
            "1": PatternFill("solid", fgColor="C6EFCE"),
            "2": PatternFill("solid", fgColor="DEEBF7"),
        }
        for rec in rows:
            sheet.append([rec.get(k, "") for k in CSV_FIELDS])
            prio = str(rec.get("приоритет") or "")
            if prio in fills:
                sheet.cell(sheet.max_row, CSV_FIELDS.index("приоритет") + 1).fill = fills[prio]
        widths = [36, 42, 18, 20, 32, 28, 28, 18, 22, 12, 28, 40]
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(CSV_FIELDS))}{max(1, len(rows)+1)}"
        sheet.freeze_panes = "A2"
    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = 22
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def stats(by_seg: dict[str, list[dict]]) -> str:
    lines = ["# Сводка базы лидов «Дед Мороз 52»", ""]
    for seg, rows in by_seg.items():
        src = Counter()
        for r in rows:
            for s in (r.get("источник") or "").split(";"):
                if s.strip():
                    src[s.strip()] += 1
        lines.append(f"## {seg}")
        lines.append(f"- записей: {len(rows)}")
        lines.append(f"- с email: {sum(bool(r.get('email')) for r in rows)}")
        lines.append(f"- с телефоном: {sum(bool(r.get('телефон')) for r in rows)}")
        lines.append(f"- с сайтом: {sum(bool(r.get('сайт')) for r in rows)}")
        lines.append(f"- с VK: {sum(bool(r.get('vk')) for r in rows)}")
        lines.append(f"- приоритет 1–2: {sum(str(r.get('приоритет')) in {'1','2'} for r in rows)}")
        lines.append("- источники: " + ", ".join(f"{k} ({v})" for k, v in src.most_common()))
        lines.append("")
    return "\n".join(lines)


def main() -> None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    index = collect_all()
    rows = collapse_bank_branches(list(index.values()))
    print(f"[build] merged unique {len(rows)}")
    missing = [r for r in rows if not r.get("email") and r.get("сайт")]
    print(f"[build] sites to enrich {len(missing)}")
    rows = enrich_rows(rows, max_workers=8, limit=500)
    by_seg = {seg: [] for seg in SEGMENTS}
    for rec in rows:
        seg = rec.get("сегмент")
        if seg in by_seg:
            by_seg[seg].append(rec)
    for seg, path in SEGMENTS.items():
        by_seg[seg] = sort_rows(by_seg[seg])
        write_csv(path, by_seg[seg])
        print(f"[build] wrote {path.name} {len(by_seg[seg])}")
    xlsx = ROOT / "leads-nn.xlsx"
    write_xlsx(by_seg, xlsx)
    report = stats(by_seg)
    (ROOT / "REPORT.md").write_text(report, encoding="utf-8")
    print(report)
    print("[build] xlsx", xlsx)


if __name__ == "__main__":
    main()
